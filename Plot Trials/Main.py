import pandas as pd
import datetime
import numpy as np
from pptx import Presentation
from pptx.util import Inches
from shapeplacement import TrialPlacement
import pptshapes as shp
import dates

start_year = 2021
end_year = 2027
start_date = datetime.datetime(start_year,1,1)
end_date = datetime.datetime(end_year,12,31)

Ms = pd.date_range(start_date, end_date, freq="MS").strftime("%Y-%m").tolist()

df = pd.read_excel(r"[LINK TO EXCEL WITH SCRAPED TRIAL INFORMATION].xlsx")

df['SSD'] = pd.to_datetime(df['SSD'], format='mixed', errors="coerce").dt.to_period('m')
df['PCD'] = pd.to_datetime(df['PCD'], format='mixed', errors="coerce").dt.to_period('m')


trials = df.to_dict()

year_list = shp.makeYearList(start_year,end_year)
year_width = 5/len(year_list)

trial_placement = TrialPlacement(Ms, 5, year_list, 4.8,start_year)
trial_placement.unit_calc()
trial_placement.month_pos()
def trialShape(trials, t, i):
    start, end = dates.getDates(trials, i)
    L, w = trial_placement.placement(start,end,i)
    t = shp.makeTrialShape(shapes, trials, L, t, w, i)
    return t

prs = Presentation(r'[INSERT LINK TO PRESENTATION TEMPLATE IF NEEDED].pptx')
title_only_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_only_slide_layout)
shapes = slide.shapes

slide_width = Inches(10)  # Adjust the width as needed
slide_height = Inches(5.63)  # Adjust the height as needed
prs.slide_width = int(slide_width)
prs.slide_height = int(slide_height)

# maketext placement
left = Inches(0.25)
width = Inches(1)
height = Inches(0.25)

trial_made = 0
top = 0
for i in range(len(trials["SSD"])):
    if i != 0 and trials["Therapy"][i] != trials["Therapy"][i - 1]:
        nextslide = prs.slides.add_slide(title_only_slide_layout)
        shapes = nextslide.shapes
        nextslide.shapes.title.text = trials["Therapy"][i] + " Trials"
        top = Inches(1.25)
    if top >= Inches(4.75) and trial_made < len(trials["SSD"]):
        nextslide = prs.slides.add_slide(title_only_slide_layout)
        shapes = nextslide.shapes
        nextslide.shapes.title.text = trials["Therapy"][i] + " Trials"
        top = Inches(1.25)
    if i == 0:
        top = Inches(1.25)
        shp.makeText(shapes,trials,left, top, width, height,i)
        if trials["KeyTrial"][i] == "Yes":
            shp.keyTrialShape(shapes,Inches(0.2), top - Inches(0.03), Inches(9.7))
        top = trialShape(trials, top, i) + Inches(0.4)
        trial_made = + 1
        slide.shapes.title.text = trials["Therapy"][i] + " Trials"
    else:
        shp.makeText(shapes,trials,left, top, width, height,i)
        if trials["KeyTrial"][i] == "Yes":
            shp.keyTrialShape(shapes,Inches(0.2), top - Inches(0.03), Inches(9.7))
        top = trialShape(trials, top, i) + Inches(0.4)
        trial_made = + 1

for j in range(len(prs.slides)):
    left_start = Inches(4.8)

    t2 = Inches(0.82)
    w2 = Inches(year_width)
    h2 = Inches(0.25)

    Lt = Inches(0.25)
    tt = Inches(2.25)
    wt = Inches(1)
    ht = Inches(0.25)

    begin_x = end_x = left_start
    end = []
    end = shp.lowestShape(prs,j) + 0.5
    end_y = Inches(end)

    shp.makeBackShape(prs,end_y, j)
    shp.makeHeaders(prs,Lt, t2, wt, ht,j)

    for year in year_list:
        shp.makeYearShape(prs,left_start, t2, w2, h2, year, j)
        left_start = left_start + Inches(year_width)

    for year in range(len(year_list) * 2 + 1):
        shp.makeLine(prs,begin_x, end_x, end_y, year,j)
        begin_x = end_x = end_x + Inches((year_width/2))

prs.save(r'[SAVE TO DESIRED LOCATION]')
